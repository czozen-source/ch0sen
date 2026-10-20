"""Builds Rejestr-uszkodzen-wozkow.xlsx: the app's tabs as an Excel workbook driven by formulas.

Sheets: Instrukcja, Flota, Zgłoszenia, Historia wózka, 1..31 (rozpiska per day of the month chosen
in Ustawienia), Pracownicy, Wózki, Ustawienia, Dane rozpisek (all day slots flattened for lookups).
Only Excel-2007-era functions, so it also recalculates in LibreOffice.

    python3 excel/build_workbook.py
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

OUT = Path(__file__).with_name("Rejestr-uszkodzen-wozkow.xlsx")

# same defaults as the app (docs/index.html)
SHIFTS = [("I", "06:00–14:00"), ("II", "14:00–22:00"), ("III", "22:00–06:00")]
SECTIONS = [("Rozładunek zwrotów", 6), ("Sprawdzanie zwrotów", 8), ("Opakowania", 4), ("Zadania dodatkowe", 4), ("Biuro", 5)]
IMPACTS = ["Uderzenie w regał", "Uderzenie w bramę / drzwi", "Uderzenie w ścianę / słup", "Kolizja z innym wózkiem",
           "Uderzenie w paletę / ładunek", "Najechanie na przeszkodę / próg", "Inne"]
PARTS = ["Widły / platforma", "Maszt / łańcuchy / hydraulika", "Hamulce", "Układ kierowniczy", "Koła / opony",
         "Oświetlenie / sygnał", "Pas / kabina", "Akumulator / ładowanie", "Obudowa / osłony", "Inne"]
WHEN = ["Zgłoszone przez sprawcę", "Zastane przy przejęciu"]
STATUSES = [("Zgłoszona", "TAK"), ("W naprawie", "TAK"), ("Czeka na części", "TAK"), ("Naprawiona", "NIE"), ("Bez naprawy", "NIE")]
YEAR, MONTH = 2026, 10

ZG_FIRST, ZG_LAST = 5, 504          # 500 zgłoszeń
DR_FIRST, DR_LAST = 2, 2600         # flattened day slots
H_FIRST, H_DMG, H_USE = 13, 100, 300
H_LAST = H_FIRST + H_DMG + 3 + H_USE - 1

# app palette
HEAD, ACCENT, LINE, MUTED = "1D2733", "E8A200", "D5DBE2", "5B6878"
OK, OK_BG, WARN, WARN_BG = "1F8A4C", "E2F4EA", "B26B00", "FFF1D6"
BUSY, BUSY_BG, BAD, BAD_BG = "1F62B8", "E1ECFA", "C0332B", "FBE4E2"
INPUT_BG, SEC_BG = "FFFBE6", "E6E6E6"

F = "Arial"
thin = Side(style="thin", color="000000")
hair = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
fill = lambda c: PatternFill("solid", start_color=c, end_color=c)
font = lambda **k: Font(name=F, **k)


def fmt_date(x):  # dd.mm.yyyy without TEXT(), whose codes are localized (Polish Excel wants "rrrr")
    return f'RIGHT("0"&DAY({x}),2)&"."&RIGHT("0"&MONTH({x}),2)&"."&YEAR({x})'


ZG = "'Zgłoszenia'"
DR = "'Dane rozpisek'"
HI = "'Historia wózka'"
zg = lambda c: f"{ZG}!${c}${ZG_FIRST}:${c}${ZG_LAST}"
dr = lambda c: f"{DR}!${c}${DR_FIRST}:${c}${DR_LAST}"

wb = Workbook()
wb.remove(wb.active)


def sheet(name, tab=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    if tab:
        ws.sheet_properties.tabColor = tab
    return ws


def title(ws, ref, text, span=None, size=16):
    ws[ref] = text
    ws[ref].font = font(bold=True, size=size, color=HEAD)
    if span:
        ws.merge_cells(f"{ref}:{span}")


def header_row(ws, row, col, labels, fill_c=HEAD, color="FFFFFF"):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t)
        c.font = font(bold=True, color=color, size=10)
        c.fill = fill(fill_c)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BOX


def inp(c):
    c.fill = fill(INPUT_BG)
    c.border = Border(left=hair, right=hair, top=hair, bottom=hair)


def dv(ws, formula, ref, warn=False, prompt=None):
    v = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True)
    if warn:
        v.errorStyle = "warning"
        v.error = "Tej wartości nie ma na liście. Zostawić mimo to?"
    if prompt:
        v.prompt, v.showInputMessage = prompt, True
    ws.add_data_validation(v)
    v.add(ref)


def cf(ws, ref, formula, **style):
    ws.conditional_formatting.add(ref, Rule(type="expression", formula=[formula], dxf=DifferentialStyle(**style)))


def widths(ws, w):
    for col, x in w.items():
        ws.column_dimensions[col].width = x


def dyn_print_area(ws, rows_formula, cols):
    q = f"'{ws.title}'"
    ws.defined_names["_xlnm.Print_Area"] = DefinedName("_xlnm.Print_Area", localSheetId=wb.sheetnames.index(ws.title),
                                                        attr_text=f"OFFSET({q}!$A$1,0,0,{rows_formula},{cols})")


def a4(ws, landscape=False, fit_width=False):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.page_margins.header = ws.page_margins.footer = 0
    if fit_width:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0


# ---------------------------------------------------------------- sheets in tab order
ws_ins = sheet("Instrukcja", ACCENT)
ws_fl = sheet("Flota", OK)
ws_zg = sheet("Zgłoszenia", WARN)
ws_hi = sheet("Historia wózka", BUSY)
days = [sheet(str(d)) for d in range(1, 32)]
ws_pr = sheet("Pracownicy", MUTED)
ws_wo = sheet("Wózki", MUTED)
ws_us = sheet("Ustawienia", MUTED)
ws_dr = sheet("Dane rozpisek", MUTED)

# ---------------------------------------------------------------- Ustawienia
ws = ws_us
title(ws, "A1", "Ustawienia")
for r, (lab, val) in enumerate([("Rok", YEAR), ("Miesiąc (1–12)", MONTH)], 3):
    ws.cell(r, 1, lab).font = font(bold=True)
    inp(ws.cell(r, 2, val))
ws["A5"], ws["B5"] = "Dni w miesiącu", "=DAY(EOMONTH(DATE(B3,B4,1),0))"
ws["A5"].font = font(bold=True)
ws["C3"] = "← wpisz rok i miesiąc: arkusze 1–31 pokażą daty tego miesiąca"
ws["C3"].font = font(italic=True, color=MUTED, size=9)
header_row(ws, 7, 1, ["Zmiana", "Godziny"])
for i, (n, h) in enumerate(SHIFTS):
    inp(ws.cell(8 + i, 1, n)); inp(ws.cell(8 + i, 2, h))
SH_ROW = 8
header_row(ws, 13, 1, ["Stanowisko", "Liczba wierszy"])
for i, (n, rows) in enumerate(SECTIONS):
    inp(ws.cell(14 + i, 1, n)); ws.cell(14 + i, 2, rows)
SEC_ROW = 14
ws["A20"] = "Nazwy zmian i stanowisk możesz zmienić; liczba wierszy na stanowisku jest stała (taka jak w aplikacji)."
ws["A20"].font = font(italic=True, color=MUTED, size=9)
lists = [("D", "Rodzaje uderzeń", IMPACTS, 30), ("F", "Co uszkodzone", PARTS, 30), ("H", "Kiedy powstało", WHEN, 10)]
for col, head, items, n in lists:
    header_row(ws, 3, ws[col + "1"].column, [head])
    for i in range(n):
        inp(ws[f"{col}{4 + i}"])
    for i, t in enumerate(items):
        ws[f"{col}{4 + i}"] = t
header_row(ws, 3, 10, ["Status", "Otwarte?"])
for i, (s, o) in enumerate(STATUSES):
    ws.cell(4 + i, 10, s); ws.cell(4 + i, 11, o)
header_row(ws, 3, 13, ["TAK / NIE"])
ws["M4"], ws["M5"] = "TAK", "NIE"
widths(ws, {"A": 24, "B": 16, "C": 4, "D": 32, "E": 3, "F": 32, "G": 3, "H": 26, "I": 3, "J": 18, "K": 10, "L": 3, "M": 10})
ST = {s: f"Ustawienia!$J${4 + i}" for i, (s, _) in enumerate(STATUSES)}

names = {
    "Rok": "Ustawienia!$B$3", "Miesiac": "Ustawienia!$B$4", "DniMiesiaca": "Ustawienia!$B$5",
    "ListaZmian": f"Ustawienia!$A${SH_ROW}:$A${SH_ROW + len(SHIFTS) - 1}",
    "ListaUderzen": "OFFSET(Ustawienia!$D$4,0,0,MAX(1,COUNTA(Ustawienia!$D$4:$D$33)),1)",
    "ListaCzesci": "OFFSET(Ustawienia!$F$4,0,0,MAX(1,COUNTA(Ustawienia!$F$4:$F$33)),1)",
    "ListaKiedy": "OFFSET(Ustawienia!$H$4,0,0,MAX(1,COUNTA(Ustawienia!$H$4:$H$13)),1)",
    "ListaStatusow": "Ustawienia!$J$4:$J$8", "TakNie": "Ustawienia!$M$4:$M$5",
    "ListaPracownikow": "OFFSET(Pracownicy!$A$2,0,0,MAX(1,COUNTA(Pracownicy!$A$2:$A$2000)),1)",
    "ListaWozkow": "OFFSET('Wózki'!$A$2,0,0,MAX(1,COUNTA('Wózki'!$A$2:$A$500)),1)",
}
for n, ref in names.items():
    wb.defined_names[n] = DefinedName(n, attr_text=ref)

# ---------------------------------------------------------------- Pracownicy / Wózki
for ws, heads, rows, n, w in [
    (ws_pr, ["Imię i nazwisko", "Dział"], [("KOWALSKI JAN", "opakowania"), ("NOWAK ANNA", "strefa"), ("WIŚNIEWSKI PIOTR", "biuro")], 2000, {"A": 34, "B": 18}),
    (ws_wo, ["Nr wózka", "Typ"], [(101, "Czołowy"), (102, "Wysoki skład"), (105, "Paleciak elektryczny"), (110, "Kompletacyjny")], 500, {"A": 12, "B": 26}),
]:
    header_row(ws, 1, 1, heads)
    for r in range(2, 2 + 200):
        inp(ws.cell(r, 1)); inp(ws.cell(r, 2))
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    widths(ws, w)
    ws.freeze_panes = "A2"
    ws["D1"] = "Przykładowe wpisy: zastąp je własną listą (można wkleić kolumny z Excela). Bez pustych wierszy w środku listy."
    ws["D1"].font = font(italic=True, color=MUTED, size=9)
dv(ws_pr, '"biuro,opakowania,strefa"', "B2:B2000", warn=True)
dv(ws_wo, '"Czołowy,Wysoki skład,Paleciak elektryczny,Kompletacyjny"', "B2:B500", warn=True)

# ---------------------------------------------------------------- 1..31: rozpiska
WEEKDAYS = '"poniedziałek","wtorek","środa","czwartek","piątek","sobota","niedziela"'
slot_rows = []      # (sheet index, shift, section, n, row)
for d, ws in enumerate(days, 1):
    widths(ws, {L(1 + 4 * s + i): x for s in range(len(SHIFTS)) for i, x in enumerate([5, 40, 12, 28])})
    for s in range(len(SHIFTS)):
        c0 = 1 + 4 * s
        A, B, C, D = (L(c0 + i) for i in range(4))
        title(ws, f"{A}1", "ROZPISKA WÓZKÓW", f"{B}1", size=14)
        ws[f"{C}1"] = f"=IF({d}<=DniMiesiaca,DATE(Rok,Miesiac,{d}),\"\")" if s == 0 else "=$C$1"
        ws[f"{C}1"].number_format = "dd.mm.yyyy"
        ws[f"{D}1"] = f'=IF({C}1="","(brak dnia {d} w tym miesiącu)",CHOOSE(WEEKDAY({C}1,2),{WEEKDAYS}))'
        for x in (C, D):
            ws[f"{x}1"].font = font(bold=True, size=12)
        ws[f"{C}1"].alignment = Alignment(horizontal="right")
        ws[f"{A}2"] = f'="ZMIANA "&Ustawienia!$A${SH_ROW + s}'
        ws[f"{C}2"] = f"=Ustawienia!$B${SH_ROW + s}&\"\""
        ws.merge_cells(f"{A}2:{B}2"); ws.merge_cells(f"{C}2:{D}2")
        for x in (A, B, C, D):
            ws[f"{x}2"].border = BOX
        ws[f"{A}2"].font = font(bold=True, size=14)
        ws[f"{C}2"].font = font(size=10)
        ws[f"{A}3"] = "Specjalista zmiany:"
        ws.merge_cells(f"{A}3:{B}3"); ws.merge_cells(f"{C}3:{D}3")
        ws[f"{A}3"].font = font(size=10)
        inp(ws[f"{C}3"])
        dv(ws, "=ListaPracownikow", f"{C}3", warn=True)
        header_row(ws, 4, c0, ["Lp", "Imię i nazwisko", "Wózek", "Uwagi"], fill_c="FFFFFF", color="000000")
        r = 5
        for si, (_, rows) in enumerate(SECTIONS):
            ws[f"{A}{r}"] = f"=UPPER(Ustawienia!$A${SEC_ROW + si})"
            ws.merge_cells(f"{A}{r}:{D}{r}")
            ws[f"{A}{r}"].font = font(bold=True, size=9.5)
            for x in (A, B, C, D):
                ws[f"{x}{r}"].fill = fill(SEC_BG); ws[f"{x}{r}"].border = BOX
            r += 1
            for n in range(rows):
                ws[f"{A}{r}"] = n + 1
                ws[f"{A}{r}"].alignment = Alignment(horizontal="center")
                for x in (A, B, C, D):
                    ws[f"{x}{r}"].border = Border(left=thin, right=thin, bottom=hair)
                    ws[f"{x}{r}"].font = font(size=10)
                for x in (B, C, D):
                    ws[f"{x}{r}"].fill = fill(INPUT_BG)
                ws[f"{C}{r}"].alignment = Alignment(horizontal="center")
                slot_rows.append((d, s, si, n, r))
                r += 1
        last = r - 1
        dv(ws, "=ListaPracownikow", f"{B}6:{B}{last}", warn=True)
        dv(ws, "=ListaWozkow", f"{C}6:{C}{last}", warn=True)
        rng = f"{C}6:{C}{last}"
        # same warnings as the app: forklift out of service, forklift twice in a shift, person twice
        cf(ws, rng, f'AND({C}6<>"",COUNTIFS({zg("B")},{C}6,{zg("L")},"TAK",{zg("O")},"TAK")>0)',
           font=Font(bold=True, color=BAD), fill=fill(BAD_BG))
        cf(ws, rng, f'AND({C}6<>"",COUNTIF(${C}$6:${C}${last},{C}6)>1)', font=Font(bold=True, color=BAD), fill=fill(BAD_BG))
        cf(ws, f"{B}6:{B}{last}", f'AND({B}6<>"",COUNTIF(${B}$6:${B}${last},{B}6)>1)', fill=fill(WARN_BG))
        ws[f"{A}{last + 2}"] = "Sporządził: ______________________"
        ws[f"{C}{last + 2}"] = "Podpis specjalisty zmiany: ________________"
        ws.merge_cells(f"{A}{last + 2}:{B}{last + 2}"); ws.merge_cells(f"{C}{last + 2}:{D}{last + 2}")
        ws[f"{A}{last + 2}"].font = ws[f"{C}{last + 2}"].font = font(size=9.5)
    for rr in range(5, last + 1):
        ws.row_dimensions[rr].height = 17
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22
    # one shift = one A4 page
    ws.print_area = f"A1:{L(4 * len(SHIFTS))}{last + 2}"
    for s in range(1, len(SHIFTS)):
        ws.col_breaks.append(Break(id=4 * s))
    a4(ws)
    ws.page_setup.scale = 92
    ws.print_options.horizontalCentered = True
DAY_LAST = last

# ---------------------------------------------------------------- Dane rozpisek
ws = ws_dr
heads = ["Data", "Zmiana", "Stanowisko", "Lp", "Imię i nazwisko", "Wózek", "Uwagi", "Klucz wózek|data",
         "Klucz historii", "Godziny", "Usterek (dzień / następny)", "Usterki (dzień / następny)", "Arkusz"]
header_row(ws, 1, 1, heads)
assert len(slot_rows) <= DR_LAST - DR_FIRST + 1
for i, (d, s, si, n, r) in enumerate(slot_rows):
    R = DR_FIRST + i
    c0 = 1 + 4 * s
    B_, C_, D_ = L(c0 + 1), L(c0 + 2), L(c0 + 3)
    sh = f"'{d}'"
    row = [
        f'=IF({sh}!$C$1="","",{sh}!$C$1)',
        f"=Ustawienia!$A${SH_ROW + s}&\"\"",
        f"=Ustawienia!$A${SEC_ROW + si}&\"\"",
        n + 1,
        f'=IF({sh}!{B_}{r}="","",TRIM({sh}!{B_}{r}&""))',
        f'=IF({sh}!{C_}{r}="","",{sh}!{C_}{r})',
        f"={sh}!{D_}{r}&\"\"",
        f'=IF(OR(A{R}="",F{R}=""),"",F{R}&"|"&A{R})',
        f'=IF(AND(H{R}<>"",OR({HI}!$B$4="",F{R}={HI}!$B$4),A{R}>={HI}!$B$5,A{R}<={HI}!$B$6),A{R}+{(9 - s) / 10}+(5000-ROW())/10000000,0)',
        f"=Ustawienia!$B${SH_ROW + s}&\"\"",
        f'=IF(H{R}="","",COUNTIFS({zg("B")},F{R},{zg("D")},">="&A{R},{zg("D")},"<="&(A{R}+1)))',
        f'=IF(OR(H{R}="",K{R}=0),"—",IF(ISNUMBER(MATCH(H{R},{zg("S")},0)),{fmt_date(f"A{R}")}&": "&INDEX({zg("H")},MATCH(H{R},{zg("S")},0)),'
        f'{fmt_date(f"(A{R}+1)")}&": "&INDEX({zg("H")},MATCH(F{R}&"|"&(A{R}+1),{zg("S")},0)))&IF(K{R}>1," (+"&(K{R}-1)&")",""))',
        d,
    ]
    for c, v in enumerate(row, 1):
        ws.cell(R, c, v).font = font(size=9)
    ws.cell(R, 1).number_format = "dd.mm.yyyy"
widths(ws, {"A": 11, "B": 8, "C": 20, "D": 5, "E": 26, "F": 8, "G": 18, "H": 14, "I": 12, "J": 12, "K": 10, "L": 30, "M": 7})
ws.freeze_panes = "A2"
ws["O1"] = "Arkusz pomocniczy: zbiera wszystkie dni z arkuszy 1–31. Nie wpisuj tu nic."
ws["O1"].font = font(italic=True, color=MUTED, size=9)

# ---------------------------------------------------------------- Zgłoszenia
ws = ws_zg
title(ws, "A1", "ZGŁOSZENIA USZKODZEŃ", "F1")
ws["A2"] = f'="Otwarte (do naprawy): "&COUNTIF({zg("O")},"TAK")&"  ·  wszystkich: "&COUNTA({zg("B")})'
ws["A2"].font = font(bold=True, color=WARN)
ws["H2"] = "Żółte kolumny wypełniasz, białe liczą się same. Filtry w nagłówku = filtry z aplikacji."
ws["H2"].font = font(italic=True, color=MUTED, size=9)
heads = ["Lp", "Wózek", "Typ wózka", "Data zdarzenia", "Zmiana", "Pracownik (operator)", "Dział", "Rodzaj uderzenia",
         "Co uszkodzone", "Kiedy powstało", "Opis zdarzenia", "Wyłączony z pracy", "Status", "Notatka / naprawa",
         "Otwarte?", "Wg rozpiski", "klucz historii", "klucz otwartych", "klucz wózek|data"]
header_row(ws, 4, 1, heads)
ws.row_dimensions[4].height = 30
for r in range(ZG_FIRST, ZG_LAST + 1):
    f = {
        1: f'=IF(B{r}="","",ROW()-{ZG_FIRST - 1})',
        3: f'=IF(B{r}="","",IFERROR(INDEX(\'Wózki\'!$B$2:$B$500,MATCH(B{r},\'Wózki\'!$A$2:$A$500,0))&"",""))',
        7: f'=IF(F{r}="","",IFERROR(INDEX(Pracownicy!$B$2:$B$2000,MATCH(F{r},Pracownicy!$A$2:$A$2000,0))&"",""))',
        15: f'=IF(B{r}="","",IF(M{r}="","TAK",IFERROR(INDEX(Ustawienia!$K$4:$K$8,MATCH(M{r},Ustawienia!$J$4:$J$8,0)),"TAK")))',
        16: f'=IF(OR(B{r}="",D{r}=""),"",IFERROR(INDEX({dr("B")},MATCH(S{r},{dr("H")},0))&": "&INDEX({dr("E")},MATCH(S{r},{dr("H")},0))'
            f'&IF(COUNTIF({dr("H")},S{r})>1," (+"&(COUNTIF({dr("H")},S{r})-1)&")",""),"—"))',
        17: f'=IF(AND(B{r}<>"",D{r}<>"",OR({HI}!$B$4="",B{r}={HI}!$B$4),D{r}>={HI}!$B$5,D{r}<={HI}!$B$6),D{r}+ROW()/1000000,0)',
        18: f'=IF(O{r}="TAK",IF(D{r}="",0,D{r})+ROW()/1000000,0)',
        19: f'=IF(OR(B{r}="",D{r}=""),"",B{r}&"|"&D{r})',
    }
    for c in range(1, 20):
        cell = ws.cell(r, c, f.get(c))
        cell.font = font(size=10)
        cell.border = Border(bottom=Side(style="thin", color=LINE))
        cell.alignment = Alignment(vertical="top", wrap_text=c in (11, 14, 16))
        if c not in f:
            cell.fill = fill(INPUT_BG)
    ws.cell(r, 4).number_format = "dd.mm.yyyy"
# example row
for c, v in {2: 105, 4: date(YEAR, MONTH, 1), 5: "I", 6: "KOWALSKI JAN", 8: IMPACTS[0], 9: PARTS[0], 10: WHEN[0],
             11: "Przykład: przy odkładaniu palety w alejce C uderzył widłami w słup regału C4; wygięty lewy ząb wideł.",
             12: "TAK", 13: "Zgłoszona"}.items():
    ws.cell(ZG_FIRST, c, v)
rng = lambda c: f"{c}{ZG_FIRST}:{c}{ZG_LAST}"
dv(ws, "=ListaWozkow", rng("B"), warn=True)
dv(ws, "=ListaZmian", rng("E"))
dv(ws, "=ListaPracownikow", rng("F"), warn=True, prompt="Zacznij pisać nazwisko — Excel 365 podpowiada z listy")
dv(ws, "=ListaUderzen", rng("H"))
dv(ws, "=ListaCzesci", rng("I"))
dv(ws, "=ListaKiedy", rng("J"))
dv(ws, "=TakNie", rng("L"))
dv(ws, "=ListaStatusow", rng("M"))
ws.add_data_validation(v := DataValidation(type="date", operator="greaterThan", formula1="36526", showErrorMessage=True, error="Wpisz datę, np. 15.10.2026"))
v.add(rng("D"))
whole = f"A{ZG_FIRST}:P{ZG_LAST}"
cf(ws, whole, f'$O{ZG_FIRST}="NIE"', font=Font(color="888888"))
st_style = [("Zgłoszona", WARN, WARN_BG), ("W naprawie", BUSY, BUSY_BG), ("Czeka na części", BUSY, BUSY_BG), ("Naprawiona", OK, OK_BG), ("Bez naprawy", MUTED, "EEEEEE")]
for s, fg, bg in st_style:
    cf(ws, rng("M"), f'$M{ZG_FIRST}="{s}"', font=Font(bold=True, color=fg), fill=fill(bg))
cf(ws, rng("L"), f'AND($L{ZG_FIRST}="TAK",$O{ZG_FIRST}="TAK")', font=Font(bold=True, color=BAD), fill=fill(BAD_BG))
widths(ws, {"A": 5, "B": 8, "C": 16, "D": 12, "E": 7, "F": 24, "G": 11, "H": 24, "I": 22, "J": 20, "K": 44, "L": 10,
            "M": 15, "N": 24, "O": 9, "P": 22, "Q": 10, "R": 10, "S": 12})
for c in "QRS":
    ws.column_dimensions[c].hidden = True
ws.freeze_panes = "C5"
ws.auto_filter.ref = f"A4:P{ZG_LAST}"
ws.print_title_rows = "4:4"
a4(ws, landscape=True, fit_width=True)
ws["S2"] = f'=MAX({ZG_FIRST},SUMPRODUCT(MAX((B{ZG_FIRST}:B{ZG_LAST}<>"")*ROW(B{ZG_FIRST}:B{ZG_LAST}))))'   # last filled row, for printing
dyn_print_area(ws, f"{ZG}!$S$2", 16)

# ---------------------------------------------------------------- Flota
ws = ws_fl
title(ws, "A1", "STAN WÓZKÓW", "E1")
ws["A2"] = "Liczone na żywo ze zgłoszeń. Lista wózków: arkusz Wózki."
ws["A2"].font = font(italic=True, color=MUTED, size=9)
FL_FIRST, FL_LAST = 8, 207
fl = lambda c: f"${c}${FL_FIRST}:${c}${FL_LAST}"
tiles = [("Sprawne", "Sprawny", OK, OK_BG), ("Usterka zgłoszona", "Usterka zgłoszona", WARN, WARN_BG),
         ("W naprawie", "W naprawie", BUSY, BUSY_BG), ("Wyłączone z pracy", "Wyłączony z pracy", BAD, BAD_BG)]
for i, (lab, val, fg, bg) in enumerate(tiles):
    c = 1 + i
    ws.cell(4, c, f'=COUNTIF({fl("C")},"{val}")').font = font(bold=True, size=20, color=fg)
    ws.cell(5, c, lab).font = font(size=9, color=fg, bold=True)
    for r in (4, 5):
        ws.cell(r, c).fill = fill(bg)
        ws.cell(r, c).alignment = Alignment(horizontal="center")
header_row(ws, 7, 1, ["Nr wózka", "Typ", "Stan", "Otwarte usterki", "Ostatnia otwarta usterka"])
for r in range(FL_FIRST, FL_LAST + 1):
    w = r - FL_FIRST + 2
    ws[f"A{r}"] = f"=IF('Wózki'!A{w}=\"\",\"\",'Wózki'!A{w})"
    ws[f"B{r}"] = f"=IF(A{r}=\"\",\"\",'Wózki'!B{w}&\"\")"
    ws[f"D{r}"] = f'=IF(A{r}="","",COUNTIFS({zg("B")},A{r},{zg("O")},"TAK"))'
    ws[f"C{r}"] = (f'=IF(A{r}="","",IF(COUNTIFS({zg("B")},A{r},{zg("O")},"TAK",{zg("L")},"TAK")>0,"Wyłączony z pracy",'
                   f'IF(COUNTIFS({zg("B")},A{r},{zg("M")},{ST["W naprawie"]})+COUNTIFS({zg("B")},A{r},{zg("M")},{ST["Czeka na części"]})>0,"W naprawie",'
                   f'IF(D{r}>0,"Usterka zgłoszona","Sprawny"))))')
    ws[f"G{r}"] = f'=IF(OR(A{r}="",D{r}=0),"",MATCH(SUMPRODUCT(MAX(({zg("B")}=A{r})*{zg("R")})),{zg("R")},0))'
    ws[f"E{r}"] = (f'=IF(A{r}="","",IF(D{r}=0,"Brak otwartych usterek",IF(INDEX({zg("D")},G{r})="","",{fmt_date(f"INDEX({zg("D")},G{r})")})'
                   f'&" · "&INDEX({zg("H")},G{r})&" — "&IF(INDEX({zg("M")},G{r})="","Zgłoszona",INDEX({zg("M")},G{r}))))')
    for c in "ABCDE":
        ws[f"{c}{r}"].font = font(size=10)
    ws[f"A{r}"].font = font(size=12, bold=True)
    ws[f"A{r}"].alignment = Alignment(horizontal="left")
    ws[f"D{r}"].alignment = Alignment(horizontal="center")
cf(ws, f"A{FL_FIRST}:E{FL_LAST}", f'$A{FL_FIRST}<>""', border=Border(bottom=Side(style="thin", color=LINE)))
for lab, val, fg, bg in tiles:
    cf(ws, f"C{FL_FIRST}:C{FL_LAST}", f'$C{FL_FIRST}="{val}"', font=Font(bold=True, color=fg), fill=fill(bg))
widths(ws, {"A": 16, "B": 22, "C": 20, "D": 16, "E": 60, "F": 3, "G": 6})
ws.column_dimensions["G"].hidden = True
ws.freeze_panes = "A8"
a4(ws, fit_width=True)
ws["G2"] = f"={FL_FIRST - 1}+MAX(1,COUNTA('Wózki'!$A$2:$A$500))"   # last filled row, for printing
dyn_print_area(ws, "'Flota'!$G$2", 5)

# ---------------------------------------------------------------- Historia wózka
ws = ws_hi
title(ws, "A1", "HISTORIA WÓZKA", "D1")
ws["A2"] = (f'=IF(B4="","Wszystkie wózki","Wózek "&B4)&" · "&IF(B5="","",{fmt_date("B5")})&" – "&IF(B6="","",{fmt_date("B6")})')
ws["A2"].font = font(bold=True, size=12)
for r, lab in [(4, "Wózek:"), (5, "Od:"), (6, "Do:")]:
    ws.cell(r, 1, lab).font = font(bold=True)
    inp(ws.cell(r, 2)); inp(ws.cell(r, 3))
    ws.merge_cells(f"B{r}:C{r}")
    ws.cell(r, 2).alignment = Alignment(horizontal="left")
ws["B5"], ws["B6"] = "=DATE(Rok,Miesiac,1)", "=EOMONTH(B5,0)"
ws["B5"].number_format = ws["B6"].number_format = "dd.mm.yyyy"
ws["D4"] = "← wybierz wózek (puste = wszystkie)"
ws["D5"] = "← domyślnie miesiąc z Ustawień; możesz wpisać dowolne daty"
for c in ("D4", "D5"):
    ws[c].font = font(italic=True, color=MUTED, size=9)
dv(ws, "=ListaWozkow", "B4", warn=True)
sums = [("Usterek w okresie", f'=COUNTIF({zg("Q")},">0")', WARN, WARN_BG),
        ("Otwartych usterek", f'=COUNTIFS({zg("Q")},">0",{zg("O")},"TAK")', BAD, BAD_BG),
        ("Dni w rozpiskach", f'=ROUND(SUMPRODUCT(({dr("I")}>0)/(COUNTIFS({dr("A")},{dr("A")},{dr("I")},">0")+({dr("I")}=0))),0)', HEAD, "F5F7F9"),
        ("Różnych operatorów", f'=ROUND(SUMPRODUCT(({dr("I")}>0)*({dr("E")}<>"")/(COUNTIFS({dr("E")},{dr("E")},{dr("I")},">0")+({dr("I")}=0))),0)', HEAD, "F5F7F9")]
for i, (lab, f, fg, bg) in enumerate(sums):
    a, b = L(1 + 2 * i), L(2 + 2 * i)
    ws[f"{a}8"], ws[f"{a}9"] = f, lab
    ws.merge_cells(f"{a}8:{b}8"); ws.merge_cells(f"{a}9:{b}9")
    ws[f"{a}8"].font = font(bold=True, size=18, color=fg)
    ws[f"{a}9"].font = font(bold=True, size=9, color=fg)
    for r in (8, 9):
        for x in (a, b):
            ws[f"{x}{r}"].fill = fill(bg); ws[f"{x}{r}"].border = Border(left=thin, right=thin, top=thin if r == 8 else None, bottom=thin if r == 9 else None)
        ws[f"{a}{r}"].alignment = Alignment(horizontal="center")
ws["A11"] = "USTERKI I USZKODZENIA"
ws["A11"].font = font(bold=True, size=12)
DMG_HEADS = ["Data", "Wózek", "Zmiana", "Uderzenie", "Co uszkodzone", "Kiedy powstało", "Operator", "Wg rozpiski", "Status", "Opis"]
USE_HEADS = ["Data", "Wózek", "Zmiana", "Operator", "Stanowisko", "Uwagi", "Usterki (dzień / następny)", "", "", ""]
header_row(ws, 12, 1, DMG_HEADS, fill_c=SEC_BG, color="000000")
# helpers: L2 = damages shown, L3 = uses shown, L5 = rows taken by the damage block, L4 = last printed row
ws["K2"], ws["L2"] = "pomocnicze:", f'=MIN(COUNTIF({zg("Q")},">0"),{H_DMG})'
ws["L3"] = f'=MIN(COUNTIF({dr("I")},">0"),{H_USE})'
ws["L5"] = "=MAX(L2,1)"
ws["L4"] = f"={H_FIRST - 1}+L5+3+L3"
EMPTY1 = "Brak zgłoszonych usterek w tym okresie."
EMPTY2 = "Brak wpisów z tym wózkiem w rozpiskach dla tego okresu."
dmg_src = [zg("D"), zg("B"), zg("E"), zg("H"), zg("I"), zg("J"), zg("F"), zg("P"), None, zg("K")]
use_src = [dr("A"), dr("F"), None, dr("E"), dr("C"), dr("G"), dr("L"), None, None, None]
for r in range(H_FIRST, H_LAST + 1):
    k = f"(ROW()-{H_FIRST - 1})"
    ws[f"L{r}"] = (f'=IF({k}<=$L$5,IF($L$2=0,"E1","D"),IF({k}=$L$5+1,"",IF({k}=$L$5+2,"T",IF({k}=$L$5+3,IF($L$3=0,"E2","H"),'
                   f'IF({k}<=$L$5+3+$L$3,"U","")))))')
    ws[f"M{r}"] = (f'=IF(L{r}="D",MATCH(LARGE({zg("Q")},{k}),{zg("Q")},0),IF(L{r}="U",MATCH(LARGE({dr("I")},{k}-$L$5-3),{dr("I")},0),""))')
    for c in range(10):
        col = L(1 + c)
        if c == 8:
            dval = f'IF(INDEX({zg("M")},$M{r})="","Zgłoszona",INDEX({zg("M")},$M{r}))'
        else:
            dval = f"INDEX({dmg_src[c]},$M{r})" + ("" if c == 0 else '&""')
        if c == 2:
            uval = f'INDEX({dr("B")},$M{r})&" · "&INDEX({dr("J")},$M{r})'
        elif use_src[c]:
            uval = f"INDEX({use_src[c]},$M{r})" + ("" if c == 0 else '&""')
        else:
            uval = '""'
        other = '""'
        if c == 0:
            other = f'IF($L{r}="T","KTO JEŹDZIŁ",IF($L{r}="E1","{EMPTY1}",IF($L{r}="E2","{EMPTY2}",IF($L{r}="H","Data",""))))'
        elif USE_HEADS[c]:
            other = f'IF($L{r}="H","{USE_HEADS[c]}","")'
        ws[f"{col}{r}"] = f'=IF($L{r}="D",{dval},IF($L{r}="U",{uval},{other}))'
        ws[f"{col}{r}"].font = font(size=9.5)
        ws[f"{col}{r}"].alignment = Alignment(vertical="top", wrap_text=c in (3, 4, 6, 7, 9))
    ws[f"A{r}"].number_format = "dd.mm.yyyy"
out = f"A{H_FIRST}:J{H_LAST}"
cf(ws, out, f'$L{H_FIRST}="T"', font=Font(bold=True, size=12, color="000000"))
use_out = f"A{H_FIRST}:G{H_LAST}"   # the "Kto jeździł" table has 7 columns
cf(ws, use_out, f'$L{H_FIRST}="H"', font=Font(bold=True), fill=fill(SEC_BG), border=BOX)
cf(ws, out, f'$L{H_FIRST}="D"', border=BOX)
cf(ws, use_out, f'$L{H_FIRST}="U"', border=BOX)
cf(ws, out, f'OR($L{H_FIRST}="E1",$L{H_FIRST}="E2")', font=Font(italic=True, color=MUTED))
cf(ws, f"G{H_FIRST}:G{H_LAST}", f'AND($L{H_FIRST}="U",G{H_FIRST}<>"—")', font=Font(bold=True, color=BAD), border=BOX)
for s, fg, bg in st_style:
    cf(ws, f"I{H_FIRST}:I{H_LAST}", f'AND($L{H_FIRST}="D",$I{H_FIRST}="{s}")', font=Font(bold=True, color=fg), fill=fill(bg), border=BOX)
widths(ws, {"A": 14, "B": 8, "C": 15, "D": 22, "E": 20, "F": 18, "G": 22, "H": 20, "I": 14, "J": 40, "K": 10, "L": 6, "M": 6})
for c in "KLM":
    ws.column_dimensions[c].hidden = True
a4(ws, landscape=True, fit_width=True)
# print area follows the filled rows, so no blank pages after the last entry
dyn_print_area(ws, f"{HI}!$L$4", 10)

# ---------------------------------------------------------------- Instrukcja
ws = ws_ins
title(ws, "A1", "REJESTR USZKODZEŃ WÓZKÓW — wersja Excel", size=18)
ws["A2"] = "Ten sam układ co aplikacja: Flota, Zgłoszenia, Rozpiska (arkusze 1–31), Historia wózka."
ws["A2"].font = font(color=MUTED)
lines = [
    ("Na start", None),
    ("1.", "Ustawienia: wpisz rok i miesiąc. Arkusze 1–31 pokażą daty i dni tygodnia tego miesiąca (w krótszych miesiącach arkusze 29–31 będą puste)."),
    ("2.", "Pracownicy i Wózki: zastąp przykładowe wpisy własną listą (można wkleić kolumny z Excela). Z tych list biorą się rozwijane listy."),
    ("3.", "Ustawienia: rodzaje uderzeń, części, nazwy zmian i stanowisk — jak w aplikacji."),
    ("Na co dzień", None),
    ("Arkusze 1–31", "Rozpiska na dany dzień: zmiana I, II i III obok siebie. Wpisz osobę, wózek i uwagi. Czerwone pole wózka = wózek wpisany dwa razy na tej zmianie albo wyłączony z pracy (otwarte zgłoszenie). Żółte nazwisko = osoba wpisana dwa razy."),
    ("Zgłoszenia", "Jeden wiersz = jedno zgłoszenie. Status zmieniasz z listy. „Wyłączony z pracy” = TAK blokuje wózek na Flocie i w rozpisce. Kolumna „Wg rozpiski” sama pokazuje, kto miał ten wózek tego dnia."),
    ("Flota", "Stan każdego wózka liczony ze zgłoszeń: Sprawny / Usterka zgłoszona / W naprawie / Wyłączony z pracy."),
    ("Historia wózka", "Wybierz wózek i okres: najpierw usterki i uszkodzenia, niżej kto jeździł (z rozpisek). Wydruk obejmuje tylko wypełnione wiersze."),
    ("Drukowanie", "Arkusz dnia: każda zmiana drukuje się na osobnej stronie A4. Plik nie ma nagłówków ani stopek z adresem."),
    ("Kolory", None),
    ("Żółte pola", "wypełniasz Ty. Białe pola liczą się same z formuł — nie nadpisuj ich."),
    ("Dane rozpisek", "arkusz pomocniczy na końcu: zbiera wpisy ze wszystkich dni. Nie wpisuj tam nic."),
    ("Nowy miesiąc", "Zapisz plik jako kopię (np. Rozpiska 11.2026.xlsx), w kopii zmień miesiąc w Ustawieniach i wyczyść arkusze 1–31."),
    ("Przykład", "W pliku jest jedno przykładowe zgłoszenie (wózek 105) i jeden wpis w rozpisce dnia 1 — usuń je przed użyciem."),
]
r = 4
for a, b in lines:
    if b is None:
        r += 1
        ws.cell(r, 1, a).font = font(bold=True, size=12, color=HEAD)
        ws.cell(r, 1).border = Border(bottom=Side(style="medium", color=ACCENT))
        ws.cell(r, 2).border = Border(bottom=Side(style="medium", color=ACCENT))
    else:
        ws.cell(r, 1, a).font = font(bold=True)
        ws.cell(r, 2, b).font = font()
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 1).alignment = Alignment(vertical="top")
    r += 1
ws["A" + str(r + 1)] = "Żółte pole"; inp(ws["A" + str(r + 1)])
ws["B" + str(r + 1)] = "– do wypełnienia"
widths(ws, {"A": 18, "B": 110})
a4(ws, fit_width=True)

# example rozpiska entry on day 1, shift I, first row of the first section
days[0]["B6"], days[0]["C6"], days[0]["D6"] = "KOWALSKI JAN", 105, "przykład"

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font and c.font.name != F:
                c.font = Font(name=F, size=c.font.size, bold=c.font.bold, italic=c.font.italic, color=c.font.color)
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print(OUT, len(slot_rows), "slotów")
